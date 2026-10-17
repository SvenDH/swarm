"""Real Monty/library integration, including snapshot restores and IO boundaries."""
import asyncio
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

from agents import Runtime, ExecutionError
from wrappers.registry import available


@unittest.skipUnless(all(name in available() for name in ('numpy', 'pandas', 'PIL', 'openpyxl', 'pypdf', 'docx')), 'optional libraries not installed')
class ImportTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.database = self.tmp.name + '/agent.db'
        self.runtime = await Runtime(self.database, workers=2).__aenter__()
        self.agent = await self.runtime.agent()

    async def asyncTearDown(self):
        await self.runtime.close()
        self.tmp.cleanup()

    async def test_numpy_imports_arrays_operators_slices_and_saved_aliases(self):
        self.assertEqual(await self.agent.execute('''
import numpy as np
from numpy import array as make_array
from numpy.linalg import norm
a = make_array([[1,2],[3,4]], dtype=float)
b = a
a += 1
a[:,0] *= 2
(a.shape, a.dtype.name, b.tolist(), float(norm([3,4])))
'''), "((2, 2), 'float64', [[4.0, 3.0], [8.0, 5.0]], 5.0)")
        self.assertEqual(await self.agent.execute('(a*2 + 1).sum()'), '44.0')
        self.assertEqual(await self.agent.execute('make_array([8,9]).tolist()'), '[8, 9]')
        self.assertEqual(await self.agent.execute('np.arange(6).reshape(2,3)[:,1:].tolist()'), '[[1, 2], [4, 5]]')
        await self.agent.execute("np.save('array.npy', a)")
        self.assertEqual(await self.agent.execute("np.load('array.npy').tolist()"), '[[4.0, 3.0], [8.0, 5.0]]')
        self.assertEqual(await self.agent.execute('original=np.arange(4); reshaped=original.reshape(2,2); reshaped[0,0]=8; original.tolist()'), '[8, 1, 2, 3]')

    async def test_pandas_construction_selection_groupby_and_file_roundtrips(self):
        self.assertEqual(await self.agent.execute('''
import pandas as pd
df = pd.DataFrame({'g':['a','a','b'], 'x':[1,2,3]})
df['y'] = df['x'] * 2
df.loc[df['x'] > 1, 'y'] = 9
df.groupby('g')['y'].sum().to_dict()
'''), "{'a': 11, 'b': 9}")
        self.assertEqual(await self.agent.execute('df.iloc[:2]["x"].tolist()'), '[1, 2]')
        self.assertEqual(await self.agent.execute('pd.Series([1,2,3]).mean()'), '2.0')
        await self.agent.execute("df.to_csv('data.csv', index=False); df.to_excel('data.xlsx', index=False)")
        self.assertEqual(await self.agent.execute("pd.read_csv('data.csv')['y'].tolist()"), '[2, 9, 9]')
        self.assertEqual(await self.agent.execute("pd.read_excel('data.xlsx')['x'].tolist()"), '[1, 2, 3]')
        self.assertEqual(await self.agent.execute('df.columns.tolist()'), "['g', 'x', 'y']")
        self.assertEqual(await self.agent.execute('df.x.sum()'), '6')
        self.assertEqual(await self.agent.execute("df.groupby('g').agg({'x':'sum'}).to_dict()"), "{'x': {'a': 3, 'b': 3}}")
        self.assertEqual(await self.agent.execute('pd.DataFrame([[1,2],[3,4]]).to_dict("records")'),
                         '[{0: 1, 1: 2}, {0: 3, 1: 4}]')
        self.assertEqual(await self.agent.execute('[row["x"] for index,row in df.iterrows()]'), '[1, 2, 3]')
        self.assertEqual(await self.agent.execute('list(df)'), "['g', 'x', 'y']")

    async def test_numpy_tile_imports_shapes_restart_and_allocation_limits(self):
        import numpy as np
        for data, reps in (([1,2], 3), ([1,2], (2,3)), ([[1],[2]], (2,3)),
                           (7, (2,3)), ([1,2], (0,2)), ([1,2], ())):
            with self.subTest(data=data, reps=reps):
                expected = np.tile(data, reps).tolist()
                self.assertEqual(await self.agent.execute(
                    f'import numpy as np; np.tile({data!r},{reps!r}).tolist()'), repr(expected))
        await self.agent.execute('from numpy import tile; tiled=tile(np.array([1,2],dtype="int32"),[2,1])')
        session_id = self.agent.id
        await self.runtime.close()
        self.runtime = await Runtime(self.database, workers=2).__aenter__()
        self.agent = await self.runtime.agent(session_id=session_id)
        self.assertEqual(await self.agent.execute('(tiled.tolist(),tiled.dtype.name)'), "([[1, 2], [1, 2]], 'int32')")
        from wrappers.numpy import tile
        with patch('wrappers.numpy.np.tile', side_effect=AssertionError('allocation reached')):
            for reps in (1_000_001, (1_000_001,0), -1, 1.5):
                with self.subTest(reps=reps), self.assertRaises((ValueError, TypeError)):
                    tile([1,2], reps)
        with self.assertRaises(ExecutionError):
            await self.agent.execute('np.tile([1,2],1000001)')

    async def test_pillow_images_mutate_persist_and_save(self):
        self.assertEqual(await self.agent.execute('''
from PIL import Image
image = Image.new('RGB', (8,8), color='red')
image.putpixel((0,0), (0,255,0))
image.save('image.png')
resized = Image.open('image.png').resize((2,2), Image.Resampling.NEAREST)
(image.size, image.getpixel((0,0)), resized.size)
'''), '((8, 8), (0, 255, 0), (2, 2))')
        self.assertEqual(await self.agent.execute('image.getpixel((0,0))'), '(0, 255, 0)')
        await self.agent.execute("resized.save('small.png')")
        self.assertTrue((await self.agent.read_file(self.agent.home + '/small.png')).startswith(b'\x89PNG'))

    async def test_extended_numpy_operations_and_durable_views(self):
        self.assertEqual(await self.agent.execute('''
import numpy as np
from numpy.linalg import cholesky
base = np.arange(6).reshape(2,3)
flipped = np.flip(base, axis=1)
flipped[0,0] = 9
(base.tolist(), np.repeat([1,2], [2,3]).tolist(), np.eye(2).tolist(),
 np.diag([2,3]).tolist(), np.cumsum([1,2,3]).tolist(),
 np.quantile([1,2,3], 0.5), cholesky([[4.,0.],[0.,9.]]).tolist())
'''), '([[0, 1, 9], [3, 4, 5]], [1, 1, 2, 2, 2], [[1.0, 0.0], [0.0, 1.0]], [[2, 0], [0, 3]], [1, 3, 6], 2.0, [[2.0, 0.0], [0.0, 3.0]])')
        self.assertEqual(await self.agent.execute('''
gx, gy = np.meshgrid([1,2],[3,4])
(gx.tolist(), gy.tolist(), np.histogram([0,1,2], bins=2)[0].tolist(),
 np.allclose(np.nan_to_num([np.nan,2]), [0,2]), base.repeat(2, axis=0).shape)
'''), '([[1, 2], [1, 2]], [[3, 3], [4, 4]], [1, 2], True, (4, 3))')
        session_id = self.agent.id
        await self.runtime.close()
        self.runtime = await Runtime(self.database, workers=2).__aenter__()
        self.agent = await self.runtime.agent(session_id=session_id)
        self.assertEqual(await self.agent.execute('flipped[1,0]=8; base.tolist()'), '[[0, 1, 9], [3, 4, 8]]')

    async def test_extended_numpy_preflights_allocations(self):
        from wrappers.numpy import repeat, eye, diag, meshgrid, histogram
        cases = [('repeat', repeat, ([1,2], 1_000_000)),
                 ('eye', eye, (2000,)), ('diag', diag, ([1], 1_000_000)),
                 ('meshgrid', meshgrid, (list(range(1001)), list(range(1001)))),
                 ('histogram', histogram, ([1,2], 10_000_000))]
        for name, function, args in cases:
            with self.subTest(name=name), patch('wrappers.numpy.np.' + name,
                    side_effect=AssertionError('native allocation reached')):
                with self.assertRaises(ValueError):
                    function(*args)

    async def test_workbook_sheet_cell_views_and_iteration(self):
        self.assertEqual(await self.agent.execute('''
from openpyxl import Workbook, load_workbook
book = Workbook()
sheet = book.active
sheet.title = 'Data'
sheet.append(['name', 'count'])
sheet.append(['a', 1])
cell = sheet.cell(2, 2, value=3)
cell.value += 1
book.save('book.xlsx')
[(row[0], row[1]) for row in sheet.iter_rows(values_only=True)]
'''), "[('name', 'count'), ('a', 4)]")
        self.assertEqual(await self.agent.execute("load_workbook('book.xlsx').active['B2'].value"), '4')
        self.assertEqual(await self.agent.execute("sheet['B2'].value=8; cell.value"), '8')
        self.assertEqual(await self.agent.execute("other = book.create_sheet('Other'); other.append([9]); book.sheetnames"), "['Data', 'Other']")
        self.assertEqual(await self.agent.execute('book.remove(other); book.sheetnames'), "['Data']")

    async def test_docx_document_paragraph_run_table_and_picture(self):
        self.assertEqual(await self.agent.execute('''
from docx import Document
from docx.shared import Inches, Pt
from PIL import Image
Image.new('RGB',(2,2),'blue').save('picture.png')
document = Document()
paragraph = document.add_paragraph('Hello')
run = paragraph.add_run(' world')
run.bold = True
run.font.size = Pt(12)
table = document.add_table(rows=2, cols=2)
table.cell(0,0).text = 'cell'
document.add_picture('picture.png', width=Inches(1))
document.save('document.docx')
(paragraph.text, run.bold, table.cell(0,0).text, len(document.inline_shapes))
'''), "('Hello world', True, 'cell', 1)")
        self.assertEqual(await self.agent.execute("Document('document.docx').paragraphs[0].text"), "'Hello world'")

    async def test_pdf_reader_writer_pages_and_copying(self):
        self.assertEqual(await self.agent.execute('''
from pypdf import PdfReader, PdfWriter
writer = PdfWriter()
writer.add_blank_page(width=100, height=100)
writer.write('source.pdf')
reader = PdfReader('source.pdf')
copy = PdfWriter()
copy.add_page(reader.pages[0])
copy.pages[0].rotate(90)
copy.write('copy.pdf')
(len(reader.pages), [page.extract_text() for page in reader.pages], len(copy.pages))
'''), "(1, [''], 1)")
        self.assertEqual(await self.agent.execute("len(PdfReader('copy.pdf').pages)"), '1')

    async def test_library_objects_survive_runtime_restart_and_rollback(self):
        await self.agent.execute('''
import numpy as np
import pandas as pd
from openpyxl import Workbook
from PIL import Image
from docx import Document
from pypdf import PdfWriter
array = np.array([1,2])
df = pd.DataFrame({'x':[3,4]})
book = Workbook()
cell = book.active['A1']
cell.value = 'before'
image = Image.new('RGB',(2,2),'red')
doc = Document()
paragraph = doc.add_paragraph('before')
writer = PdfWriter()
writer.add_blank_page(width=10,height=10)
''')
        session_id = self.agent.id
        await self.runtime.close()
        self.runtime = await Runtime(self.database, workers=2).__aenter__()
        self.agent = await self.runtime.agent(session_id=session_id)
        self.assertEqual(await self.agent.execute('(array.sum(),df["x"].sum(),cell.value,paragraph.text,image.size,len(writer.pages))'),
                         "(3, 7, 'before', 'before', (2, 2), 1)")
        with self.assertRaises(ExecutionError):
            await self.agent.execute("array[0]=99; cell.value='bad'; paragraph.text='bad'; raise ValueError('rollback')")
        self.assertEqual(await self.agent.execute('(array.tolist(),cell.value,paragraph.text)'), "([1, 2], 'before', 'before')")
        await self.agent.execute("cell.value='after'; book.save('after.xlsx'); paragraph.text='after'; doc.save('after.docx')")

    async def test_standard_python_still_works_after_library_imports(self):
        self.assertEqual(await self.agent.execute('''
import numpy as np, math
from pathlib import Path
calls = []
def index():
    calls.append(1)
    return 0
items = [1,2]
alias = items
items += [3]
items[index()] += 4
Path('standard.txt').write_text('ok')
(items,alias,len(calls),math.sqrt(9),[x+1 for x in range(3)],Path('standard.txt').read_text())
'''), "([5, 2, 3], [5, 2, 3], 1, 3.0, [1, 2, 3], 'ok')")

    async def test_unsupported_and_host_file_apis_are_not_available(self):
        from PIL import Image as NativeImage
        from pypdf import PdfWriter as NativeWriter
        secret = Path(self.tmp.name) / 'host-secret.txt'
        secret.write_text('secret')
        host_image = Path(self.tmp.name) / 'host.png'
        NativeImage.new('RGB', (2, 2)).save(host_image)
        host_pdf = Path(self.tmp.name) / 'host.pdf'
        writer = NativeWriter()
        writer.add_blank_page(width=10, height=10)
        writer.write(host_pdf)
        snippets = [
            "import numpy as np; np.__dict__",
            "import pandas as pd; pd.read_pickle('anything')",
            "import pandas as pd; pd.DataFrame({'x':[1]}).agg('to_pickle', 'anything')",
            f"import pandas as pd; pd.read_csv({str(secret)!r})",
            f"from PIL import Image; Image.open({str(secret)!r})",
            f"from PIL import Image; from pathlib import Path; Image.open(Path({str(host_image)!r}))",
            f"from pypdf import PdfWriter; PdfWriter(clone_from={str(secret)!r})",
            f"from pypdf import PdfWriter; from pathlib import Path; PdfWriter(clone_from=Path({str(host_pdf)!r}))",
            f"from pypdf import PdfWriter; from pathlib import Path; PdfWriter().append(Path({str(host_pdf)!r}))",
            "import numpy as np; np.load('anything', allow_pickle=True)",
            "from numpy import *",
            "import numpy as np; np.zeros((1000000000,))",
            "import numpy as np; np.array([1],dtype='V999999999')",
            "import numpy as np; array=np.array([1,2]); np.clip(array,0,1,out=array)",
        ]
        for snippet in snippets:
            with self.subTest(snippet=snippet), self.assertRaises(ExecutionError):
                await self.agent.execute(snippet)
        self.assertEqual(secret.read_text(), 'secret')
        self.assertIsNone((await self.agent.state()).pending)
        await self.agent.execute(f"from pathlib import Path; Path({str(secret.parent)!r}).mkdir(parents=True,exist_ok=True); import pandas as pd; pd.DataFrame({{'x':[1]}}).to_csv({str(secret)!r},index=False)")
        self.assertEqual(await self.agent.read_file(str(secret)), b'x\n1\n')
        self.assertEqual(secret.read_text(), 'secret')

    async def test_sessions_do_not_share_mutable_library_objects(self):
        a, b = self.agent, await self.runtime.agent()
        await asyncio.gather(a.execute('import numpy as np; x=np.array([1,2])'),
                             b.execute('import numpy as np; x=np.array([3,4])'))
        await a.execute('x[0]=9')
        self.assertEqual(await a.execute('x.tolist()'), '[9, 2]')
        self.assertEqual(await b.execute('x.tolist()'), '[3, 4]')

    async def test_random_generator_state_is_checkpointed_not_shared_host_state(self):
        import numpy as np
        expected = np.random.default_rng(123)
        first = expected.integers(0, 100, size=4).tolist()
        second = expected.integers(0, 100, size=4).tolist()
        self.assertEqual(await self.agent.execute('import numpy as np; rng=np.random.default_rng(123); rng.integers(0,100,size=4).tolist()'), repr(first))
        session_id = self.agent.id
        await self.runtime.close()
        self.runtime = await Runtime(self.database, workers=2).__aenter__()
        self.agent = await self.runtime.agent(session_id=session_id)
        with self.assertRaises(ExecutionError):
            await self.agent.execute("rng.integers(0,100,size=4); raise ValueError('rollback')")
        self.assertEqual(await self.agent.execute('rng.integers(0,100,size=4).tolist()'), repr(second))
