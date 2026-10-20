"""Expanded APIs exercise real Monty, native adapters, persistence and rejection."""
import tempfile
import unittest
from unittest.mock import Mock

from agents import Runtime, ExecutionError
from wrappers.registry import available


class ExpandedApiTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.database = self.tmp.name + '/agent.db'
        self.runtime = await Runtime(self.database, workers=2).__aenter__()
        self.agent = await self.runtime.agent()

    async def asyncTearDown(self):
        await self.runtime.close()
        self.tmp.cleanup()

    def need(self, *names):
        if not all(name in available() for name in names):
            self.skipTest('optional host packages not installed')

    async def reopen(self):
        session_id = self.agent.id
        await self.agent.close()
        self.agent = await self.runtime.agent(session_id=session_id)

    async def reject(self, code):
        with self.assertRaises(ExecutionError):
            await self.agent.execute(code)

    async def test_numeric_functions_machine_info_and_rng(self):
        self.need('numpy')
        self.assertEqual(await self.agent.execute('''
import numpy as np
rng=np.random.default_rng(5)
first=rng.exponential(size=3)
info=np.iinfo(np.uint8)
(np.pad([1,2],(1,2)).tolist(),np.take([[1,2],[3,4]],[1],axis=1).tolist(),
 info.max, np.finfo(np.float32).bits, np.sinh([0]).tolist(),np.union1d([1,2],[2,3]).tolist())
'''), '([0, 1, 2, 0, 0], [[2], [4]], 255, 32, [0.0], [1, 2, 3])')
        await self.reopen()
        import numpy as np
        native = np.random.default_rng(5)
        native.exponential(size=3)
        self.assertEqual(await self.agent.execute('rng.exponential(size=3).tolist()'), repr(native.exponential(size=3).tolist()))
        self.assertEqual(await self.agent.execute('info.min'), '0')
        await self.reject('np.pad([1], (1,1), mode="empty")')
        await self.reject('np.get_include()')
        await self.reject('np.frompyfunc(lambda x:x,1,1)')

    async def test_pillow_fonts_colors_chops_and_constants(self):
        self.need('PIL')
        self.assertEqual(await self.agent.execute('''
from PIL import Image, ImageFont, ImageDraw, ImageColor, ImageChops, ImageOps, ImageFilter
im=Image.new('RGB',(80,40),'navy')
font=ImageFont.load_default(size=12)
draw=ImageDraw.Draw(im)
draw.text((2,2),'Hi',font=font,fill='white')
effect=ImageFilter.UnsharpMask(radius=1,percent=120)
other=ImageChops.invert(im).transpose(Image.Transpose.FLIP_LEFT_RIGHT).filter(effect)
(ImageColor.getrgb('rebeccapurple'),other.size,ImageOps.scale(im,0.5).size,
 Image.Transform.AFFINE,Image.Dither.NONE)
'''), '((102, 51, 153), (80, 40), (40, 20), 0, 0)')
        await self.reopen()
        self.assertEqual(await self.agent.execute("draw.text((30,2),'ok',font=font); font.size"), '12')
        self.assertEqual(await self.agent.execute('im.filter(effect).filter(ImageFilter.RankFilter(3,4)).size'), '(80, 40)')
        await self.reject("ImageFont.truetype('/etc/passwd',12)")
        await self.reject('Image.eval(im,lambda x:x)')

    async def test_scipy_constants_special_distributions_and_images(self):
        self.need('scipy')
        self.assertEqual(await self.agent.execute('''
from scipy import constants, special, ndimage, stats
dist=stats.gamma(2)
(constants.c,constants.physical_constants['speed of light in vacuum'][0],
 special.jv(0,0),ndimage.zoom([[1.,2.],[3.,4.]],2,order=0).shape,
 float(dist.mean()),float(stats.bernoulli(0.25).pmf(1)))
'''), '(299792458.0, 299792458.0, 1.0, (4, 4), 2.0, 0.25)')
        await self.reopen()
        self.assertEqual(await self.agent.execute('float(dist.mean())'), '2.0')
        await self.reject('ndimage.generic_filter([[1]],lambda x:1,size=1)')
        await self.reject('special.jv(0,[0],out=[0])')
        await self.reject('special.jv(0,[0],where=[False])')
        await self.reject('special.jv(0,[0],[0])')
        await self.reject('stats.norm.expect(lambda x:x)')
        await self.reject("constants.physical_constants['speed of light in vacuum']=(1.,'m/s',0.)")

    async def test_dataframe_accessors_reducers_and_categories(self):
        self.need('pandas')
        self.assertEqual(await self.agent.execute('''
import pandas as pd
df=pd.DataFrame({'key':['a','a','b'],'value':[1,2,3]})
dates=pd.Series(pd.to_datetime(['2024-01-02','2025-03-04']))
cat=pd.Categorical(['b','a','b'])
(df['key'].str.upper().tolist(),dates.dt.year.tolist(),
 df.pivot_table(index='key',values='value',aggfunc='sum')['value'].tolist(),cat.codes.tolist())
'''), "(['A', 'A', 'B'], [2024, 2025], [3, 3], [1, 0, 1])")
        await self.reopen()
        self.assertEqual(await self.agent.execute('cat.tolist()'), "['b', 'a', 'b']")
        for reducer in ('to_pickle', 'eval', 'to_csv'):
            await self.reject(f"pd.pivot_table(df,index='key',values='value',aggfunc={reducer!r})")
        await self.reject('df.apply(lambda x:x)')

    async def test_office_enums_styles_and_roundtrip(self):
        self.need('openpyxl', 'docx', 'pptx')
        self.assertEqual(await self.agent.execute('''
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from pptx.enum.shapes import MSO_SHAPE
book=Workbook()
cell=book.active['A1']
cell.value='hello'
cell.font=Font(bold=True,color='FF0000')
cell.fill=PatternFill(fill_type='solid',fgColor='00FF00')
cell.alignment=Alignment(horizontal='center')
book.save('styled.xlsx')
(cell.font.bold,cell.alignment.horizontal,WD_ALIGN_PARAGRAPH.JUSTIFY,WD_BREAK.PAGE,MSO_SHAPE.STAR_5_POINT)
'''), "(True, 'center', 3, 7, 92)")
        await self.reopen()
        self.assertEqual(await self.agent.execute('cell.font.bold'), 'True')
        await self.reject("book.active.add_image('/etc/passwd')")

    async def test_colormap_palettes_and_recorded_plotting(self):
        self.need('matplotlib', 'seaborn')
        self.assertEqual(await self.agent.execute('''
import matplotlib.pyplot as plt
import seaborn as sns
from matplotlib.colors import CSS4_COLORS
from matplotlib import cm
fig,ax=plt.subplots(figsize=(2,2))
ax.set_facecolor('ivory')
ax.set_box_aspect(1)
sns.ecdfplot(x=[1,2,3],ax=ax)
sns.despine(ax=ax)
fig.set_facecolor('white')
fig.savefig('ecdf.png')
(CSS4_COLORS['rebeccapurple'],len(sns.color_palette('deep',3)),cm.viridis([0.,1.]).shape)
'''), "('#663399', 3, (2, 4))")
        await self.reject('sns.set_theme()')
        await self.reject("sns.color_palette('viridis',as_cmap=True)")

    async def test_symbolic_trees_and_pdf_transform(self):
        self.need('sympy', 'pypdf')
        self.assertEqual(await self.agent.execute('''
import sympy as sp
from pypdf import Transformation
x=sp.Symbol('x')
expr=sp.sinh(x)+sp.gamma(x)
solutions=sp.linsolve([x-2],[x])
transform=Transformation().translate(2,3)
(str(expr),str(solutions),transform.apply_on((1,1)))
'''), "('sinh(x) + gamma(x)', '{(2,)}', (3.0, 4.0))")
        await self.reopen()
        self.assertEqual(await self.agent.execute('str(expr)'), "'sinh(x) + gamma(x)'")
        await self.reject("sp.sympify('__import__(\"os\")')")

    async def test_reporting_constants_and_virtual_output(self):
        self.need('reportlab', 'xlsxwriter', 'pymupdf')
        self.assertEqual(await self.agent.execute('''
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.pagesizes import A5
from reportlab.lib.colors import HexColor, orange
from xlsxwriter.utility import xl_rowcol_to_cell
import pymupdf
c=Canvas('report.pdf',pagesize=A5)
c.setFillColor(HexColor('#ff8000'))
c.roundRect(10,10,40,40,4,fill=1)
c.setDash([2,2])
c.line(0,0,80,80)
c.save()
(xl_rowcol_to_cell(1,2),pymupdf.TEXT_PRESERVE_WHITESPACE,orange)
'''), "('C2', 2, (1.0, 0.6470588235294118, 0.0))")

    async def test_new_allocation_guards_run_before_native_calls(self):
        self.need('numpy', 'scipy', 'sympy')
        import numpy as np
        from wrappers.scipy import square_factory, convolution, window, image_filter
        native = Mock(side_effect=AssertionError('native allocation reached'))
        for function, args in ((square_factory(native), (10000000,)),
                               (window(native), (1000000000,)),
                               (convolution(native), (np.zeros((10001,1)),np.zeros((1,10001))))):
            with self.subTest(args=str(args)[:60]), self.assertRaises(MemoryError):
                function(*args)
        native.assert_not_called()
        def gaussian(input, sigma, radius=None, truncate=4.):
            raise AssertionError('native filter allocation reached')
        with self.assertRaises(MemoryError):
            image_filter(gaussian)([1], 10**12)
        with self.assertRaises(MemoryError):
            image_filter(gaussian)([1], 1, radius=10**12)
        await self.reject('import sympy as sp; sp.zeros(10000000,10000000)')
