"""Real optional-library smoke tests, durable state and host-availability gates."""
import asyncio
import tempfile
import unittest
import subprocess
import sys
from pathlib import Path
from unittest.mock import patch

from agents import Runtime, ExecutionError
from wrappers import tools
from wrappers import registry
from wrappers.imports import prepare


class AvailabilityTests(unittest.TestCase):
    def test_missing_packages_are_not_advertised_registered_or_importable(self):
        registry.available.cache_clear()
        try:
            with patch.object(registry, 'find_spec', return_value=None):
                self.assertEqual(registry.available(), {})
                self.assertEqual(tools(None), {})
                self.assertNotIn('numpy:', registry.help_text())
                for source in ('import numpy', 'import matplotlib.pyplot as plt',
                               'from bs4 import BeautifulSoup', 'import fitz'):
                    with self.subTest(source=source), self.assertRaisesRegex(ImportError, 'host package'):
                        prepare(source)
        finally:
            registry.available.cache_clear()

    def test_discovery_checks_dependencies_without_importing_them(self):
        registry.available.cache_clear()
        try:
            with patch.object(registry, 'find_spec', side_effect=lambda name: None if name == 'matplotlib' else object()):
                self.assertNotIn('matplotlib', registry.available())
                self.assertNotIn('seaborn', registry.available())
                self.assertIn('bs4', registry.available())
        finally:
            registry.available.cache_clear()


class OptionalTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.database = self.tmp.name + '/agent.db'
        self.runtime = await Runtime(self.database, workers=2).__aenter__()
        self.agent = await self.runtime.agent()

    async def asyncTearDown(self):
        await self.runtime.close()
        self.tmp.cleanup()

    def need(self, *names):
        if not all(name in registry.available() for name in names):
            self.skipTest('Optional packages not installed: ' + ', '.join(names))

    async def test_runtime_operates_without_any_optional_packages(self):
        registry.available.cache_clear()
        try:
            with patch.object(registry, 'find_spec', return_value=None):
                self.assertEqual(await self.agent.execute("open('plain.txt','w').write('works'); open('plain.txt').read()"), "'works'")
                with self.assertRaisesRegex(ExecutionError, 'host package'):
                    await self.agent.execute('import numpy')
                self.assertEqual(await self.agent.execute('2+3'), '5')
        finally:
            registry.available.cache_clear()

    async def restart(self):
        session_id = self.agent.id
        await self.runtime.close()
        self.runtime = await Runtime(self.database, workers=2).__aenter__()
        self.agent = await self.runtime.agent(session_id=session_id)

    async def test_expanded_plotting_recipes_and_scipy(self):
        self.need('numpy', 'matplotlib', 'scipy')
        self.assertEqual(await self.agent.execute('''
import numpy as np
import matplotlib.pyplot as plt
from scipy.special import softmax, expit
from scipy.stats import rankdata, iqr
fig, axes = plt.subplots(2, 2, figsize=(4,3))
axes[0][0].step([0,1,2], [1,3,2])
axes[0][0].annotate('peak', xy=(1,3), xytext=(0,2), arrowprops={'arrowstyle':'->'})
axes[0][0].axhspan(0,1,alpha=0.2)
axes[0][1].stairs([2,1,3])
axes[0][1].tick_params(axis='x', labelrotation=20)
axes[1][0].contourf([[1,2],[3,4]])
axes[1][0].set_aspect('equal')
axes[1][1].stackplot([0,1,2], [1,2,3], [2,1,2])
axes[1][1].margins(x=0)
fig.supxlabel('shared x')
fig.supylabel('shared y')
fig.subplots_adjust(hspace=0.5, wspace=0.5)
fig.text(0.5,0.98,'extended',ha='center')
fig.savefig('extended.png')
(np.round(softmax([0.,0.]), 2).tolist(), float(expit(0.)), rankdata([3,1,1]).tolist(), float(iqr([1,2,3,4])))
'''), '([0.5, 0.5], 0.5, [3.0, 1.5, 1.5], 1.5)')
        before = await self.agent.read_file(self.agent.home + '/extended.png')
        self.assertTrue(before.startswith(b'\x89PNG'))
        await self.restart()
        await self.agent.execute("fig.savefig('restored.png'); fig.savefig('restored.svg')")
        self.assertEqual(before, await self.agent.read_file(self.agent.home + '/restored.png'))
        for code in ("axes[0][0].text(0,0,'unsafe',usetex=True)",
                     "axes[0][0].annotate('unsafe',xy=(0,0),fontproperties={'fname':'/tmp/host.ttf'})"):
            with self.subTest(code=code), self.assertRaises(ExecutionError):
                await self.agent.execute(code)

    async def test_matplotlib_and_seaborn(self):
        self.need('matplotlib', 'seaborn')
        await self.agent.execute('''
import matplotlib.pyplot as plt
import seaborn as sns
fig, ax = plt.subplots(figsize=(3,2))
ax.plot([1,2,3],[2,4,3],label='line')
ax.set(title='Sample', xlabel='x')
ax=sns.scatterplot(x=[1,2], y=[3,4], ax=ax)
ax.set_ylabel('returned axes')
ax.legend()
fig.savefig('plot.png')
''')
        first = await self.agent.read_file(self.agent.home + '/plot.png')
        self.assertTrue(first.startswith(b'\x89PNG'))
        await self.restart()
        await self.agent.execute("fig.savefig('again.png'); plt.savefig('pyplot.svg')")
        self.assertEqual(first, await self.agent.read_file(self.agent.home + '/again.png'))
        self.assertIn(b'<svg', await self.agent.read_file(self.agent.home + '/pyplot.svg'))
        await self.agent.execute("sns.histplot(x=[1,2,2,3]).figure.savefig('hist.pdf')")
        self.assertTrue((await self.agent.read_file(self.agent.home + '/hist.pdf')).startswith(b'%PDF'))

    async def test_scipy_stats_signal_fft_and_interpolation(self):
        self.need('scipy')
        self.assertEqual(await self.agent.execute('''
from scipy import stats, signal, fft
from scipy.interpolate import interp1d
result = stats.ttest_1samp([1,2,3,4,5],3)
f = interp1d([0,1,2], [0,10,20])
(result.statistic, result.pvalue, f([0.5,1.5]).tolist(), signal.find_peaks([0,2,0])[0].tolist(), fft.fft([1,0]).real.tolist())
'''), '(0.0, 1.0, [5.0, 15.0], [1], [1.0, 1.0])')
        await self.restart()
        self.assertEqual(await self.agent.execute('(f([1.0]).tolist(), stats.norm.cdf(0))'), '([10.0], 0.5)')
        self.assertEqual(await self.agent.execute('normal=stats.norm(loc=2,scale=3); normal.mean()'), '2.0')
        self.assertEqual(await self.agent.execute('''
from scipy.optimize import linprog
solution=linprog([-1,-2],A_ub=[[1,1]],b_ub=[2],bounds=[(0,None),(0,None)])
(solution.success, solution.x.tolist(), len(stats.norm.rvs(size=3,random_state=1)))
'''), '(True, [0.0, 2.0], 3)')

    async def test_scipy_fft_checkpoint_reopens_in_a_fresh_host_process(self):
        self.need('scipy')
        await self.agent.execute('from scipy import fft; fft.fft([1,0]).tolist()')
        session_id = self.agent.id
        await self.runtime.close()
        program = '''
import asyncio, sys
from agents import Runtime
async def run():
    async with Runtime(sys.argv[1],workers=1) as runtime:
        agent=await runtime.agent(session_id=int(sys.argv[2]))
        print(await agent.execute('fft.fft([1,0]).real.tolist()'))
asyncio.run(run())
'''
        result = await asyncio.to_thread(subprocess.run,
            [sys.executable, '-X', 'faulthandler', '-c', program, self.database, str(session_id)],
            cwd=Path(__file__).resolve().parents[1], capture_output=True, text=True, timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(result.stdout.strip(), '[1.0, 1.0]')

    async def test_output_paths_are_virtual_even_when_matching_existing_host_files(self):
        self.need('matplotlib', 'reportlab', 'pptx', 'xlsxwriter')
        host = Path(self.tmp.name) / 'host-secret.txt'
        host.write_bytes(b'secret')
        await self.agent.execute(f"from pathlib import Path; Path({str(host.parent)!r}).mkdir(parents=True,exist_ok=True)")
        await self.agent.execute(f'''
import matplotlib.pyplot as plt
from pptx import Presentation
from reportlab.pdfgen.canvas import Canvas
import xlsxwriter
plt.plot([1,2]); plt.savefig({str(host)!r},format='png')
Presentation().save({str(host)!r})
c=Canvas({str(host)!r}); c.drawString(1,1,'safe'); c.save()
book=xlsxwriter.Workbook({str(host)!r}); book.add_worksheet().write('A1','safe'); book.close()
''')
        self.assertEqual(host.read_bytes(), b'secret')
        self.assertTrue((await self.agent.read_file(str(host))).startswith(b'PK'))

    async def test_presentation_text_tables_pictures_and_restart(self):
        self.need('pptx', 'PIL')
        await self.agent.execute('''
from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image
Image.new('RGB',(2,2),'red').save('pic.png')
deck=Presentation()
slide=deck.slides.add_slide(deck.slide_layouts[1])
slide.shapes.title.text='Hello'
box=slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
box.text='Body'
paragraph=box.text_frame.add_paragraph()
run=paragraph.add_run()
run.text='bold'
run.font.bold=True
table=slide.shapes.add_table(2,2,0,0,Inches(2),Inches(1)).table
table.cell(0,0).text='cell'
slide.shapes.add_picture('pic.png',0,0,width=Inches(1))
deck.save('deck.pptx')
''')
        self.assertTrue((await self.agent.read_file(self.agent.home + '/deck.pptx')).startswith(b'PK'))
        await self.restart()
        self.assertEqual(await self.agent.execute("(run.text,run.font.bold,table.cell(0,0).text,Presentation('deck.pptx').slides[0].shapes.title.text)"),
                         "('bold', True, 'cell', 'Hello')")

    async def test_reportlab_pdfplumber_and_pymupdf(self):
        self.need('reportlab', 'pdfplumber', 'pymupdf')
        await self.agent.execute('''
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
c=canvas.Canvas('report.pdf',pagesize=A4)
c.setFont('Helvetica',12)
c.drawString(72,720,'Hello PDF')
''')
        await self.restart()
        self.assertEqual(await self.agent.execute('''
c.save()
import pdfplumber
import pymupdf as fitz
pdf=pdfplumber.open('report.pdf')
doc=fitz.open('report.pdf')
page=doc[0]
page.insert_text((72,100),'Added')
page.get_pixmap(dpi=72).save('page.png')
doc.save('changed.pdf')
(pdf.pages[0].extract_text(), doc.page_count, 'Added' in page.get_text())
'''), "('Hello PDF', 1, True)")
        await self.restart()
        self.assertEqual(await self.agent.execute("('Added' in page.get_text(), pdf.pages[0].crop((0,0,500,800)).extract_text())"), "(True, 'Hello PDF')")
        self.assertTrue((await self.agent.read_file(self.agent.home + '/page.png')).startswith(b'\x89PNG'))
        self.assertEqual(await self.agent.execute("import fitz; len(fitz.open('changed.pdf'))"), '1')

    async def test_beautifulsoup_select_and_mutable_tag_views(self):
        self.need('bs4')
        self.assertEqual(await self.agent.execute('''
from bs4 import BeautifulSoup
soup=BeautifulSoup('<html><p class="x">Hello</p><a href="/url">Link</a></html>', 'html.parser')
tag=soup.select_one('p.x')
tag.string='Changed'
(soup.get_text(), soup.find('a')['href'], len(soup.select('p')))
'''), "('ChangedLink', '/url', 1)")
        await self.restart()
        self.assertEqual(await self.agent.execute("tag['id']='ok'; soup.find('p')['id']"), "'ok'")

    async def test_sympy_expressions_solve_matrix_and_restart(self):
        self.need('sympy')
        self.assertEqual(await self.agent.execute('''
import sympy as sp
x=sp.Symbol('x')
expression=(x+1)**2
str(sp.expand(expression))
'''), "'x**2 + 2*x + 1'")
        await self.restart()
        self.assertEqual(await self.agent.execute("(str(sp.diff(expression,x)), str(sp.solve(sp.Eq(x+1,3),x)[0]), str(sp.Matrix([[1,2],[3,4]]).det()))"),
                         "('2*x + 2', '2', '-2')")
        self.assertEqual(await self.agent.execute('(float(sp.Rational(1,2)), int(sp.Integer(3)), str(sp.pi))'), "(0.5, 3, 'pi')")
        self.assertEqual(await self.agent.execute('y=sp.Symbol("y"); solved=sp.solve([sp.Eq(x+y,3),sp.Eq(x-y,1)],[x,y]); int(solved[x])'), '2')

    async def test_sklearn_fit_transform_predict_survive_restart(self):
        self.need('sklearn')
        self.assertEqual(await self.agent.execute('''
from sklearn.preprocessing import StandardScaler, OneHotEncoder, LabelEncoder
from sklearn.linear_model import LinearRegression, LogisticRegression, Ridge
from sklearn.cluster import KMeans
from sklearn.decomposition import PCA
from sklearn.metrics import accuracy_score
scaler=StandardScaler()
scaled=scaler.fit_transform([[0.0],[2.0]])
model=LinearRegression().fit([[0],[1],[2]],[1,3,5])
encoder=OneHotEncoder()
encoder.fit([['a'],['b']])
labels=LabelEncoder()
labels.fit(['a','b'])
classifier=LogisticRegression().fit([[0],[1],[8],[9]],[0,0,1,1])
kmeans=KMeans(n_clusters=2,n_init=1).fit([[0,0],[0,1],[8,8],[8,9]])
pca=PCA(n_components=1).fit([[0,0],[1,1],[2,2]])
(scaled.tolist(), model.predict([[3]]).round(6).tolist(), accuracy_score([1,0],[1,0]))
'''), '([[-1.0], [1.0]], [7.0], 1.0)')
        await self.restart()
        self.assertEqual(await self.agent.execute("(model.predict([[4]]).round(6).tolist(), scaler.transform([[1]]).tolist(), encoder.transform([['b']]).tolist(), labels.transform(['b']).tolist(), classifier.predict([[10]]).tolist(), len(kmeans.predict([[0,0]])), pca.transform([[3,3]]).shape)"),
                         '([9.0], [[0.0]], [[0.0, 1.0]], [1], [1], 1, (1, 1))')
        self.assertEqual(await self.agent.execute('model.coef_.round(6).tolist()'), '[2.0]')

    async def test_xlsxwriter_sheets_formats_and_charts(self):
        self.need('xlsxwriter', 'openpyxl')
        await self.agent.execute('''
import xlsxwriter
book=xlsxwriter.Workbook('generated.xlsx')
sheet=book.add_worksheet('Data')
bold=book.add_format({'bold':True})
sheet.write_row('A1',['x','y'],bold)
sheet.write_row('A2',[1,2])
sheet.write_row('A3',[2,4])
chart=book.add_chart({'type':'line'})
chart.add_series({'categories':'=Data!$A$2:$A$3','values':'=Data!$B$2:$B$3'})
sheet.insert_chart('D1',chart)
''')
        await self.restart()
        self.assertEqual(await self.agent.execute('''
sheet.write_formula('C2','=A2+B2')
book.close()
from openpyxl import load_workbook
loaded=load_workbook('generated.xlsx')
(loaded.active['B3'].value,loaded.active['C2'].value)
'''), "(4, '=A2+B2')")

    async def test_unsafe_apis_fail_and_host_files_are_unchanged(self):
        self.need('matplotlib', 'sympy', 'reportlab', 'pymupdf', 'bs4')
        host = Path(self.tmp.name) / 'secret.txt'
        host.write_text('secret')
        snippets = [
            f"import pymupdf; pymupdf.open({str(host)!r})",
            f"import pdfplumber; pdfplumber.open({str(host)!r})",
            f"from pptx import Presentation; Presentation({str(host)!r})",
            "import sympy; sympy.sympify('__import__(\"os\").getcwd()')",
            "import sympy; sympy.simplify('__import__(\"os\").getcwd()')",
            "from bs4 import BeautifulSoup; BeautifulSoup('<xml/>','xml')",
            "import matplotlib.pyplot as plt; plt.rcParams",
            "import matplotlib.pyplot as plt; plt.plot([1],usetex=True)",
            "from reportlab.pdfgen.canvas import Canvas; c=Canvas('x.pdf'); c.setFont('/host.ttf',12)",
            "import pymupdf; doc=pymupdf.open(); page=doc.new_page(); page.insert_text((1,1),'x',fontfile='/host.ttf')",
            f"from scipy.optimize import linprog; linprog([1], options={{'log_file':{str(host)!r}}})",
        ]
        for snippet in snippets:
            with self.subTest(snippet=snippet), self.assertRaises(ExecutionError):
                await self.agent.execute(snippet)
        self.assertEqual(host.read_text(), 'secret')

    async def test_plotting_isolation_and_rollback(self):
        self.need('matplotlib')
        other = await self.runtime.agent()
        await asyncio.gather(self.agent.execute("import matplotlib.pyplot as plt; plt.plot([1,2],[3,4])"),
                             other.execute("import matplotlib.pyplot as plt; plt.plot([1,2],[8,2])"))
        await asyncio.gather(self.agent.execute("plt.savefig('a.png')"), other.execute("plt.savefig('b.png')"))
        before = await self.agent.read_file(self.agent.home + '/a.png')
        self.assertNotEqual(before, await other.read_file(other.home + '/b.png'))
        with self.assertRaises(ExecutionError):
            await self.agent.execute("plt.plot([1,2],[99,99]); plt.savefig('bad.png'); raise ValueError('rollback')")
        await self.agent.execute("plt.savefig('after.png')")
        self.assertEqual(before, await self.agent.read_file(self.agent.home + '/after.png'))


if __name__ == '__main__':
    unittest.main()
