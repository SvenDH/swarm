from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl import styles
from openpyxl.styles.proxy import StyleProxy
from openpyxl.utils import get_column_letter, column_index_from_string, range_boundaries, quote_sheetname
from copy import copy

from .files import source, save

STYLES = {'Font': styles.Font, 'PatternFill': styles.PatternFill, 'Border': styles.Border,
          'Side': styles.Side, 'Alignment': styles.Alignment, 'Protection': styles.Protection,
          'Color': styles.Color}


def exports(fs):
    def load(filename, **kwargs):
        if set(kwargs) - {'data_only'}:
            raise ValueError('Only data_only is supported for load_workbook')
        return load_workbook(source(fs, filename), keep_links=False, **kwargs)
    return {'Workbook': Workbook, 'load_workbook': load,
            'styles': {**STYLES, 'numbers': {name: value for name, value in vars(styles.numbers).items()
                                           if name.startswith('FORMAT_') and type(value) is str}},
            'utils': {'get_column_letter': get_column_letter,
                      'column_index_from_string': column_index_from_string,
                      'range_boundaries': range_boundaries, 'quote_sheetname': quote_sheetname}}


def members(value, fs, name=None):
    if isinstance(value, StyleProxy):
        value = copy(value)
    if isinstance(value, styles.Font):
        return {'name': value.name, 'size': value.sz, 'sz': value.sz, 'bold': value.b, 'b': value.b,
                'italic': value.i, 'i': value.i, 'underline': value.u, 'color': value.color}
    if isinstance(value, styles.Alignment):
        return {'horizontal': value.horizontal, 'vertical': value.vertical, 'wrap_text': value.wrap_text,
                'text_rotation': value.text_rotation, 'shrink_to_fit': value.shrink_to_fit, 'indent': value.indent}
    if isinstance(value, styles.PatternFill):
        return {'fill_type': value.fill_type, 'fgColor': value.fgColor, 'bgColor': value.bgColor}
    if isinstance(value, styles.Side):
        return {'style': value.style, 'color': value.color}
    if isinstance(value, styles.Border):
        return {'left': value.left, 'right': value.right, 'top': value.top, 'bottom': value.bottom}
    if isinstance(value, styles.Color):
        return {'type': value.type, 'rgb': value.rgb if value.type == 'rgb' else None,
                'indexed': value.indexed if value.type == 'indexed' else None, 'tint': value.tint}
    if isinstance(value, styles.Protection):
        return {'locked': value.locked, 'hidden': value.hidden}
    if isinstance(value, Workbook):
        return {'active': value.active, 'sheetnames': value.sheetnames, 'worksheets': value.worksheets,
                'create_sheet': value.create_sheet, 'remove': value.remove,
                'save': lambda filename: save(fs, filename, value.save), 'close': lambda: None}
    if isinstance(value, Worksheet):
        def rows(*args, **kwargs):
            if (kwargs.get('max_row', value.max_row) - kwargs.get('min_row', 1) + 1) * (
                    kwargs.get('max_col', value.max_column) - kwargs.get('min_col', 1) + 1) > 100_000:
                raise ValueError('iter_rows is limited to 100,000 cells')
            return list(value.iter_rows(*args, **kwargs))
        return {'title': value.title, 'max_row': value.max_row, 'max_column': value.max_column,
                'append': value.append, 'cell': value.cell, 'iter_rows': rows,
                'values': Values(value), 'merge_cells': value.merge_cells,
                'unmerge_cells': value.unmerge_cells, 'delete_rows': value.delete_rows,
                'delete_cols': value.delete_cols}
    if isinstance(value, (Cell, MergedCell)):
        return {'value': value.value, 'row': value.row, 'column': value.column,
                'font': value.font, 'fill': value.fill, 'border': value.border,
                'alignment': value.alignment, 'protection': value.protection,
                'coordinate': value.coordinate, 'number_format': value.number_format,
                'data_type': value.data_type}
    return {}


def is_view(value):
    return isinstance(value, (Worksheet, Cell, MergedCell))


class Values:
    def __init__(self, sheet):
        self.sheet = sheet

    def __iter__(self):
        if self.sheet.max_row * self.sheet.max_column > 100_000:
            raise ValueError('worksheet iteration is limited to 100,000 cells')
        return iter(self.sheet.values)


def iterate(value):
    if isinstance(value, Worksheet) and value.max_row * value.max_column > 100_000:
        raise ValueError('worksheet iteration is limited to 100,000 cells')
    return iter(value)


def assign(value, name, item):
    if isinstance(value, Cell) and name == 'value':
        value.value = item
    elif isinstance(value, Cell) and name == 'number_format':
        value.number_format = item
    elif isinstance(value, Cell) and name in ('font', 'fill', 'border', 'alignment', 'protection'):
        if not isinstance(item, tuple(STYLES.values())):
            raise TypeError('use an openpyxl.styles value')
        if name == 'font':
            value.font = item
        elif name == 'fill':
            value.fill = item
        elif name == 'border':
            value.border = item
        elif name == 'alignment':
            value.alignment = item
        else:
            value.protection = item
    elif isinstance(value, Worksheet) and name == 'title':
        value.title = item
    else:
        raise AttributeError('Unsupported workbook attribute assignment')
