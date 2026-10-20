from docx import Document
from docx.document import Document as DocumentType
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from docx.text.font import Font
from docx.table import Table, _Cell, _Row
from docx.shared import Inches, Cm, Mm, Pt, Emu, Twips, RGBColor
from docx.enum import text, table, section, style, shape
from .api import enums
from docx.shape import InlineShape
import io

from .files import source, save
from .bridge import descriptor


def exports(fs):
    return {'Document': lambda docx=None: Document() if docx is None else Document(source(fs, docx)),
            'shared': {'Inches': Inches, 'Cm': Cm, 'Mm': Mm, 'Pt': Pt, 'Emu': Emu, 'Twips': Twips,
                       'RGBColor': lambda r, g, b: tuple(RGBColor(r, g, b))},
            'enum': {name: enums(module) for name, module in (
                ('text', text), ('table', table), ('section', section), ('style', style), ('shape', shape))}}


def members(value, fs, name=None):
    if isinstance(value, DocumentType):
        return {'paragraphs': value.paragraphs, 'tables': value.tables, 'inline_shapes': value.inline_shapes,
                'add_paragraph': value.add_paragraph, 'add_heading': value.add_heading,
                'add_table': value.add_table, 'add_page_break': value.add_page_break,
                'add_picture': lambda image_path_or_stream, **kwargs:
                    value.add_picture(source(fs, image_path_or_stream), **kwargs),
                'save': lambda path_or_stream: save(fs, path_or_stream, value.save)}
    if isinstance(value, Paragraph):
        return {'text': value.text, 'runs': value.runs, 'add_run': value.add_run,
                'alignment': value.alignment}
    if isinstance(value, Table):
        return {'rows': value.rows, 'columns': value.columns, 'cell': value.cell,
                'add_row': value.add_row, 'add_column': value.add_column}
    if isinstance(value, _Row):
        return {'cells': value.cells}
    if isinstance(value, _Cell):
        return {'text': value.text, 'paragraphs': value.paragraphs,
                'add_paragraph': value.add_paragraph}
    if isinstance(value, Run):
        return {'text': value.text, 'bold': value.bold, 'italic': value.italic,
                'underline': value.underline, 'font': value.font}
    if isinstance(value, Font):
        return {'name': value.name, 'size': value.size, 'bold': value.bold, 'italic': value.italic}
    if isinstance(value, InlineShape):
        return {'width': value.width, 'height': value.height}
    return {}


def pack(value):
    if isinstance(value, DocumentType):
        output = io.BytesIO()
        value.save(output)
        return descriptor('docx', 'document', output.getvalue())


def restore(kind, data):
    if kind != 'document':
        raise ValueError('Unsupported document object')
    return Document(io.BytesIO(data))


def view(root, value):
    if not isinstance(root, DocumentType):
        return None
    if isinstance(value, InlineShape):
        for i, shape in enumerate(root.inline_shapes):
            if shape._inline is value._inline:
                return [['attr', 'inline_shapes'], ['item', i]]
    for i, paragraph in enumerate(root.paragraphs):
        path = [['attr', 'paragraphs'], ['item', i]]
        if isinstance(value, Paragraph) and value._p is paragraph._p:
            return path
        if isinstance(value, Run):
            for j, run in enumerate(paragraph.runs):
                if value._r is run._r:
                    return [*path, ['attr', 'runs'], ['item', j]]
    for i, table in enumerate(root.tables):
        path = [['attr', 'tables'], ['item', i]]
        if isinstance(value, Table) and value._tbl is table._tbl:
            return path
        for j, row in enumerate(table.rows):
            row_path = [*path, ['attr', 'rows'], ['item', j]]
            if isinstance(value, _Row) and value._tr is row._tr:
                return row_path
            if isinstance(value, _Cell):
                for k, cell in enumerate(row.cells):
                    if value._tc is cell._tc:
                        return [*row_path, ['attr', 'cells'], ['item', k]]


def is_view(value):
    return isinstance(value, (Paragraph, Run, Table, _Cell, _Row, InlineShape))


def iterate(value):
    return iter(value)


def assign(value, name, item):
    if isinstance(value, (Paragraph, Run, _Cell)) and name == 'text':
        value.text = item
    elif isinstance(value, (Paragraph, Table)) and name == 'style':
        value.style = item
    elif isinstance(value, Paragraph) and name == 'alignment':
        value.alignment = item
    elif isinstance(value, (Run, Font)) and name == 'bold':
        value.bold = item
    elif isinstance(value, (Run, Font)) and name == 'italic':
        value.italic = item
    elif isinstance(value, Run) and name == 'underline':
        value.underline = item
    elif isinstance(value, Font) and name == 'name':
        value.name = item
    elif isinstance(value, Font) and name == 'size':
        value.size = item
    elif isinstance(value, InlineShape) and name == 'width':
        value.width = item
    elif isinstance(value, InlineShape) and name == 'height':
        value.height = item
    else:
        raise AttributeError('Unsupported document attribute assignment')
