from pptx import Presentation
from pptx.presentation import Presentation as PresentationType
from pptx.slide import Slide, SlideLayout, Slides, SlideLayouts
from pptx.shapes.shapetree import SlideShapes
from pptx.shapes.base import BaseShape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.text.text import TextFrame, _Paragraph, _Run, Font
from pptx.table import Table, _Cell
from pptx.util import Inches, Cm, Mm, Pt, Emu, Centipoints
from pptx.enum import text, shapes, chart, dml, action, lang
from .api import enums
from pptx.dml.color import RGBColor, ColorFormat
from .files import source, save
from .common import iterate


def exports(fs):
    return {'Presentation': lambda path=None: Presentation(None if path is None else source(fs, path)),
            'util': {'Inches': Inches, 'Cm': Cm, 'Mm': Mm, 'Pt': Pt, 'Emu': Emu, 'Centipoints': Centipoints},
            'dml': {'color': {'RGBColor': lambda r, g, b: tuple(RGBColor(r, g, b))}},
            'enum': {name: enums(module) for name, module in (
                ('text', text), ('shapes', shapes), ('chart', chart), ('dml', dml),
                ('action', action), ('lang', lang))}}


def members(value, fs, name=None):
    if isinstance(value, PresentationType):
        return {'slides': value.slides, 'slide_layouts': value.slide_layouts,
                'slide_width': value.slide_width, 'slide_height': value.slide_height,
                'save': lambda path: save(fs, path, value.save)}
    if isinstance(value, (Slide, SlideLayout)):
        return {'shapes': value.shapes, 'placeholders': value.placeholders, 'name': value.name}
    if isinstance(value, Slides):
        return {'add_slide': value.add_slide}
    if isinstance(value, SlideShapes):
        return {'title': value.title, 'add_textbox': value.add_textbox,
                'add_shape': value.add_shape, 'add_table': value.add_table,
                'add_picture': lambda image_file, *a, **kw: value.add_picture(source(fs, image_file), *a, **kw)}
    if isinstance(value, BaseShape):
        result = {'name': value.name, 'left': value.left, 'top': value.top,
                  'width': value.width, 'height': value.height, 'has_text_frame': value.has_text_frame}
        if value.has_text_frame:
            result.update(text=value.text, text_frame=value.text_frame)
        if isinstance(value, GraphicFrame) and value.has_table:
            result['table'] = value.table
        return result
    if isinstance(value, TextFrame):
        return {'text': value.text, 'paragraphs': value.paragraphs,
                'add_paragraph': value.add_paragraph, 'clear': value.clear, 'word_wrap': value.word_wrap}
    if isinstance(value, _Paragraph):
        return {'text': value.text, 'runs': value.runs, 'font': value.font,
                'add_run': value.add_run, 'level': value.level, 'alignment': value.alignment}
    if isinstance(value, _Run):
        return {'text': value.text, 'font': value.font}
    if isinstance(value, Font):
        return {'name': value.name, 'size': value.size, 'bold': value.bold,
                'italic': value.italic, 'color': value.color}
    if isinstance(value, ColorFormat):
        return {'rgb': None if value.type is None else tuple(value.rgb)}
    if isinstance(value, Table):
        return {'cell': value.cell}
    if isinstance(value, _Cell):
        return {'text': value.text, 'text_frame': value.text_frame}
    return {}


def text_view(frame, value, path):
    if isinstance(value, TextFrame) and value._txBody is frame._txBody:
        return path
    for i, paragraph in enumerate(frame.paragraphs):
        prefix = [*path, ['attr', 'paragraphs'], ['item', i]]
        if isinstance(value, _Paragraph) and value._p is paragraph._p:
            return prefix
        for j, run in enumerate(paragraph.runs):
            if isinstance(value, _Run) and value._r is run._r:
                return [*prefix, ['attr', 'runs'], ['item', j]]


def is_view(value):
    return isinstance(value, (Slide, SlideLayout, Slides, SlideLayouts, SlideShapes,
                              BaseShape, TextFrame, _Paragraph, _Run, Font, Table, _Cell, ColorFormat))


def assign(value, name, item):
    if isinstance(value, (BaseShape, TextFrame, _Paragraph, _Run, _Cell)) and name == 'text':
        value.text = item
    elif isinstance(value, Font) and name == 'size':
        value.size = item
    elif isinstance(value, Font) and name == 'bold':
        value.bold = item
    elif isinstance(value, Font) and name == 'italic':
        value.italic = item
    elif isinstance(value, Font) and name == 'name':
        value.name = item
    elif isinstance(value, ColorFormat) and name == 'rgb':
        value.rgb = RGBColor(*item)
    elif isinstance(value, _Paragraph) and name == 'alignment':
        value.alignment = item
    elif isinstance(value, _Paragraph) and name == 'level':
        value.level = item
    elif isinstance(value, TextFrame) and name == 'word_wrap':
        value.word_wrap = item
    elif isinstance(value, PresentationType) and name == 'slide_width':
        value.slide_width = item
    elif isinstance(value, PresentationType) and name == 'slide_height':
        value.slide_height = item
    else:
        raise AttributeError('Unsupported presentation attribute')
