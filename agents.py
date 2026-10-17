"""Monty code agent: events + immutable objects are the durable source of truth.

sessions/files and visible workspaces are rebuildable views. Changes go through
agent tools; editing a workspace directly does not change recorded state.
execute() resumes pending work; execution_id deduplicates retries. rebuild()
replays events without running code. Inject Storage to replace local files.
The database and Monty snapshots must be trusted. Library wrappers expose plain
data and virtual files; their host work is outside Monty's resource accounting.
"""
from __future__ import annotations

import fnmatch
import hashlib
import io
import json
import os
import posixpath
import re
import shutil
import sqlite3
import uuid
from pathlib import Path as HostPath
from pathlib import PurePosixPath as Path
from typing import Protocol
from contextlib import contextmanager
from dataclasses import dataclass

import httpx
import numpy as np
import pandas as pd
from PIL import Image
from docx import Document
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader, PdfWriter
from pydantic_monty import CallbackFile, Monty, MontyError, MontyCrashedError, OSAccess
from dotenv import load_dotenv
from names import num_to_name

load_dotenv(override=True)


def code_prefix(arguments):
    """Decode the complete characters of an unfinished JSON code argument."""
    start = re.search(r'"code"\s*:\s*"', arguments)
    if start is None:
        return ''
    token = re.match(r'(?:[^"\\\x00-\x1f]|\\(?:["\\/bfnrt]|u[0-9a-fA-F]{4}))*',
                     arguments[start.end():]).group()
    code = json.loads('"' + token + '"')
    return code[:-1] if code and 0xD800 <= ord(code[-1]) <= 0xDBFF else code


def sse_data(response):
    """SSE events can span multiple data lines and arbitrary network chunks."""
    data = []
    for line in response.iter_lines():
        if line.startswith('data:'):
            data.append(line[5:].lstrip(' '))
        elif not line and data:
            yield '\n'.join(data)
            data.clear()
    if data:
        yield '\n'.join(data)

HELP = """You are a code-mode assistant. Execute Python with the python tool to
inspect and modify a persistent virtual Unix filesystem. No host shell/network.
Globals persist after successful calls. Failed calls roll back globals and files.
Use print for observations. Batch related work; avoid dumping large files.
Paths are POSIX; cwd starts at /home/<your-name>/. pathlib.Path and open use
the same virtual filesystem as the tools. /tmp and /dev/null are available.
HOME, USER, LOGNAME, PWD and TMPDIR describe your isolated virtual environment.
cd() changes tool paths immediately and native open/Path cwd on the NEXT call.
Available host functions (call directly, no imports):
  pwd(); cd(path); ls(path='.'); mkdir(path, parents=True)
  read(path, offset=1, limit=200) -> numbered lines (1-based)
  write(path, content: str|bytes); edit(path, old, new, all=False)
  glob(pattern='*', path='/', limit=200) -> recursive matching file paths
  grep(pattern, path='/', glob='*', limit=100, offset=0, literal=False)
    -> {matches: [{path,line,text}], truncated}; literal searches avoid regex
  mv(source,target); rm(path, recursive=False)
  numpy(op, data, other=None, axis=None): array, sum, mean, min, max, std,
    median, transpose, reshape(other=shape), dot(other=array), unique
    Arrays cross the boundary as lists; scalars as Python numbers.
  pandas(op, data=None, path=None, columns=None): read_csv/read_excel -> records;
    to_csv/to_excel write records to path; describe/head (columns=count, default 5),
    select(columns=list), sort(columns=list), drop_duplicates -> plain data
  PIL(op, path=None, output=None, size=None, mode='RGB', color='white'):
    info, new, resize, convert. Writing requires output; format from extension.
  openpyxl(op, path, rows=None, sheet=None): read -> {sheet: rows},
    write -> new workbook, append -> existing workbook (sheet defaults active)
  pypdf(op, path, output=None, pages=None): read -> page texts,
    select -> write selected zero-based pages (default all)
  docx(op, path, paragraphs=None): read -> paragraphs, write -> new document,
    append -> existing document. All paths refer ONLY to virtual files.
These are functional wrappers, not importable native modules or native objects.
Finish with a concise answer based on actual execution results.
"""


class Unix(OSAccess):
    """Monty's in-memory Unix tree plus bounded code-agent conveniences."""

    def __init__(self, entries=(), cwd='/', environ=None):
        # Like Bashkit, the runtime and tools share a VFS, including virtual
        # devices. This callback never opens a host path.
        super().__init__(files=[CallbackFile('/dev/null', read=lambda path: b'',
                                             write=lambda path, data: None)], environ=environ)
        self.cwd = cwd
        for path, content in sorted(entries, key=lambda row: (row[0].count('/'), row[0])):
            if path == '/':
                continue
            p = Path(path)
            if content is None:
                self.path_mkdir(p, True, True)
            else:
                self.path_mkdir(p.parent, True, True)
                self.path_write_bytes(p, bytes(content))

    def path(self, path):
        return Path(posixpath.normpath('/' + posixpath.join(self.cwd, str(path)).lstrip('/')))

    def walk(self, path='/'):
        p = self.path(path)
        yield p
        if self.path_is_dir(p):
            for child in sorted(self.path_iterdir(p)):
                yield from self.walk(str(child))

    def entries(self):
        return [(str(p), self.path_read_bytes(p) if self.path_is_file(p) else None)
                for p in self.walk()]

    def pwd(self):
        return self.cwd

    def cd(self, path):
        p = self.path(path)
        self.path_iterdir(p)  # validates directory
        self.cwd = str(p)
        self.environ['PWD'] = self.cwd
        return self.cwd

    def ls(self, path='.'):
        return [str(p) + ('/' if self.path_is_dir(p) else '')
                for p in sorted(self.path_iterdir(self.path(path)))][:1000]

    def mkdir(self, path, parents=True):
        self.path_mkdir(self.path(path), parents, True)

    def read(self, path, offset=1, limit=200):
        if offset < 1 or not 1 <= limit <= 2000:
            raise ValueError('offset >= 1 and 1 <= limit <= 2000 required')
        lines = self.path_read_text(self.path(path)).splitlines()
        return '\n'.join(f'{i + 1}: {line[:4000]}' for i, line in
                         enumerate(lines[offset-1:offset-1+limit], offset-1))

    def write(self, path, content):
        if isinstance(content, str):
            content = content.encode()
        if len(content) > 32 * 1024 * 1024:
            raise ValueError('file exceeds 32 MiB')
        return self.path_write_bytes(self.path(path), content)

    def edit(self, path, old, new, all=False):
        text = self.path_read_text(self.path(path))
        count = text.count(old)
        if not old or not count or (count != 1 and not all):
            raise ValueError('old must be nonempty and match once, or use all=True')
        self.write(path, text.replace(old, new, -1 if all else 1))
        return count

    def glob(self, pattern='*', path='/', limit=200):
        if not 1 <= limit <= 2000:
            raise ValueError('limit must be 1..2000')
        root = self.path(path)
        result = []
        for p in self.walk(str(root)):
            if self.path_is_file(p) and (fnmatch.fnmatch(str(p.relative_to(root)), pattern)
                                         or fnmatch.fnmatch(p.name, pattern)):
                result.append(str(p))
                if len(result) == limit:
                    break
        return result

    def grep(self, pattern, path='/', glob='*', limit=100, offset=0, literal=False):
        # Host regex needs its own timeout: Monty's budget does not cover callbacks.
        import regex
        if not 1 <= limit <= 1000 or offset < 0:
            raise ValueError('limit must be 1..1000; offset >= 0')
        matcher = None if literal else regex.compile(pattern)
        result, skipped = [], 0
        for p in self.walk(path):
            if not self.path_is_file(p) or not fnmatch.fnmatch(p.name, glob):
                continue
            data = self.path_read_bytes(p)
            if b'\x00' in data:
                continue
            for number, line in enumerate(data.decode('utf-8', 'replace').splitlines(), 1):
                found = pattern in line if literal else matcher.search(line, timeout=0.05)
                if found:
                    if skipped < offset:
                        skipped += 1
                    elif len(result) == limit:
                        return {'matches': result, 'truncated': True}
                    else:
                        result.append({'path': str(p), 'line': number, 'text': line[:4000]})
        return {'matches': result, 'truncated': False}

    def mv(self, source, target):
        src, dst = self.path(source), self.path(target)
        if src == Path('/') or src == dst or src in dst.parents:
            raise ValueError('invalid move')
        self.path_rename(src, dst)
        if self.cwd == str(src) or src in Path(self.cwd).parents:
            self.cwd = str(dst / Path(self.cwd).relative_to(src))

    def path_rename(self, path, target):
        # Monty 1.1 OSAccess moves a file's tree entry but leaves its name stale.
        # Keep its metadata aligned so subsequent unlink/rename uses the new name.
        super().path_rename(path, target)
        if self.path_is_file(target):
            self._get_file(target).path = Path(target)
            self._get_file(target).name = Path(target).name

    def rm(self, path, recursive=False):
        p = self.path(path)
        if p == Path('/') or p == Path(self.cwd) or p in Path(self.cwd).parents:
            raise ValueError('cannot remove root or cwd ancestor')
        if self.path_is_dir(p):
            if recursive:
                for child in self.path_iterdir(p):
                    self.rm(str(child), True)
            self.path_rmdir(p)
        else:
            self.path_unlink(p)


    def source(self, path):
        return io.BytesIO(self.path_read_bytes(self.path(path)))

    def save(self, path, writer):
        output = io.BytesIO()
        writer(output)
        return self.write(path, output.getvalue())

    def numpy(self, op, data, other=None, axis=None):
        a = np.asarray(data)
        operations = {'array': lambda: a, 'sum': lambda: a.sum(axis=axis),
                      'mean': lambda: a.mean(axis=axis), 'min': lambda: a.min(axis=axis),
                      'max': lambda: a.max(axis=axis), 'std': lambda: a.std(axis=axis),
                      'median': lambda: np.median(a, axis=axis),
                      'transpose': lambda: a.T, 'reshape': lambda: a.reshape(other),
                      'dot': lambda: a.dot(other), 'unique': lambda: np.unique(a)}
        value = operations[op]()
        return value.tolist()

    def pandas(self, op, data=None, path=None, columns=None):
        if op in ('read_csv', 'read_excel'):
            frame = (pd.read_csv(self.source(path)) if op == 'read_csv'
                     else pd.read_excel(self.source(path)))
        else:
            frame = pd.DataFrame(data)
            if op in ('to_csv', 'to_excel'):
                writer = frame.to_csv if op == 'to_csv' else frame.to_excel
                return self.save(path, lambda output: writer(output, index=False))
            operations = {'head': lambda: frame.head(5 if columns is None else columns),
                          'describe': frame.describe, 'select': lambda: frame[columns],
                          'sort': lambda: frame.sort_values(columns),
                          'drop_duplicates': frame.drop_duplicates}
            frame = operations[op]()
        return json.loads(frame.to_json(orient='records'))

    def PIL(self, op, path=None, output=None, size=None, mode='RGB', color='white'):
        if op not in ('info', 'new', 'resize', 'convert'):
            raise ValueError('unknown image operation')
        image = Image.new(mode, tuple(size), color) if op == 'new' else Image.open(self.source(path))
        try:
            if op == 'info':
                return {'size': list(image.size), 'mode': image.mode, 'format': image.format}
            if op == 'resize':
                image = image.resize(tuple(size))
            elif op == 'convert':
                image = image.convert(mode)
            fmt = Image.registered_extensions()[Path(output).suffix.lower()]
            return self.save(output, lambda stream: image.save(stream, format=fmt))
        finally:
            image.close()

    def openpyxl(self, op, path, rows=None, sheet=None):
        if op not in ('read', 'write', 'append'):
            raise ValueError('unknown workbook operation')
        book = Workbook() if op == 'write' else load_workbook(self.source(path))
        try:
            if op == 'read':
                return {s.title: [list(row) for row in s.values] for s in book.worksheets}
            target = book[sheet] if sheet in book.sheetnames else book.active
            if sheet and sheet not in book.sheetnames:
                target = book.create_sheet(sheet)
            for row in rows or []:
                target.append(row)
            return self.save(path, book.save)
        finally:
            book.close()

    def pypdf(self, op, path, output=None, pages=None):
        reader = PdfReader(self.source(path))
        indices = range(len(reader.pages)) if pages is None else pages
        if op == 'read':
            return [reader.pages[i].extract_text() for i in indices]
        if op != 'select':
            raise ValueError('unknown PDF operation')
        writer = PdfWriter()
        for i in indices:
            writer.add_page(reader.pages[i])
        return self.save(output, writer.write)

    def docx(self, op, path, paragraphs=None):
        if op not in ('read', 'write', 'append'):
            raise ValueError('unknown document operation')
        document = Document() if op == 'write' else Document(self.source(path))
        if op == 'read':
            return [p.text for p in document.paragraphs]
        for text in paragraphs or []:
            document.add_paragraph(text)
        return self.save(path, document.save)


def sync_directory(path):
    descriptor = os.open(path, os.O_RDONLY | os.O_DIRECTORY)
    try:
        os.fsync(descriptor)
    finally:
        os.close(descriptor)


class Storage(Protocol):
    """put must durably store immutable bytes; checkout is a rebuildable view."""
    def create(self, name: str) -> None: ...
    def put(self, name: str, content: bytes) -> str: ...
    def read(self, key: str) -> bytes: ...
    def checkout(self, name: str, files: list[dict]) -> HostPath | str: ...


class DirectoryStorage:
    """Durable objects plus disposable workspace directories; replace for blobs."""

    def __init__(self, root):
        self.root = HostPath(root)
        self._mkdir(self.root.absolute())
        self.root = self.root.resolve()

    def _mkdir(self, path):
        if not path.exists():
            self._mkdir(path.parent)
            path.mkdir(exist_ok=True)
            sync_directory(path.parent)

    def _path(self, name):
        if name in ('', '.', '..') or Path(name).name != name:
            raise ValueError('invalid workspace name')
        path = self.root / name
        if path.is_symlink():
            raise PermissionError('workspace symlinks are not supported')
        return path

    def create(self, name):
        if self._path(name).exists():
            raise FileExistsError(name)

    def put(self, name, content):
        self._path(name)
        key = f'.objects/{name}/{hashlib.sha256(content).hexdigest()}'
        target = self.root / key
        self._mkdir(target.parent)
        if target.exists():
            self.read(key)
        else:
            temporary = target.with_suffix('.' + uuid.uuid4().hex)
            with temporary.open('xb') as file:
                file.write(content)
                file.flush()
                os.fsync(file.fileno())
            temporary.replace(target)
            sync_directory(target.parent)
        return key

    def read(self, key):
        parts = Path(key).parts
        if (len(parts) != 3 or parts[0] != '.objects' or parts[1] in ('.', '..')
                or len(parts[2]) != 64 or any(c not in '0123456789abcdef' for c in parts[2])):
            raise PermissionError('invalid storage key')
        host = self.root.joinpath(*parts)
        if host.resolve() != host:
            raise PermissionError('storage symlinks are not supported')
        content = host.read_bytes()
        if hashlib.sha256(content).hexdigest() != parts[2]:
            raise OSError('stored object failed its content hash check')
        return content

    def checkout(self, name, files):
        root = self._path(name)
        if root.exists():
            shutil.rmtree(root)
        root.mkdir()
        for file in files:
            if '..' in Path(file['path']).parts:
                raise ValueError('invalid virtual path')
            target = root / file['path'].lstrip('/')
            if file['kind'] == 'directory':
                target.mkdir(parents=True, exist_ok=True)
            else:
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(self.read(file['storage_key']))
        return root


@dataclass
class Session:
    id: int
    name: str
    workspace: str
    cwd: str
    monty: bytes | None
    messages: list[dict]
    generation: str
    pending: dict | None
    last_event_id: int


class ExecutionError(RuntimeError):
    """A recorded sandbox failure; infrastructure failures remain retryable."""


class CodeAgent:
    """Commit request, run code, then atomically record objects/checkpoint/result.

    Retry pending requests; return recorded results for completed execution IDs.
    Model requests may repeat after a crash. Committed local effects do not.
    """

    def __init__(self, database='agent.db', *, session_id=None, workspaces=None, storage: Storage | None = None,
                 model=None, base_url=None, api_key=None, on_output=None):
        self.on_output = on_output
        self._streamed_calls = set()
        if storage is not None and workspaces is not None:
            raise ValueError('pass storage or workspaces, not both')
        self.storage = storage if storage is not None else DirectoryStorage(
            workspaces if workspaces is not None else HostPath(database).absolute().parent / 'workspaces')
        self.db = sqlite3.connect(database, timeout=30, isolation_level=None)
        self.db.execute('PRAGMA journal_mode=WAL')
        self.db.execute('PRAGMA synchronous=FULL')
        columns = [r[1] for r in self.db.execute('PRAGMA table_info(sessions)')]
        if columns and 'last_event_id' not in columns:
            self.db.close()
            raise RuntimeError('unsupported database schema; use a fresh database and workspace root')
        self.db.executescript('''
            CREATE TABLE IF NOT EXISTS events (
                id INTEGER PRIMARY KEY AUTOINCREMENT, session_id INTEGER NOT NULL,
                kind TEXT NOT NULL, execution_id TEXT, payload TEXT NOT NULL,
                monty BLOB, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(session_id,execution_id,kind));
            CREATE INDEX IF NOT EXISTS events_session ON events(session_id,id);
            CREATE TRIGGER IF NOT EXISTS events_no_update BEFORE UPDATE ON events
                BEGIN SELECT RAISE(ABORT,'events are append-only'); END;
            CREATE TRIGGER IF NOT EXISTS events_no_delete BEFORE DELETE ON events
                BEGIN SELECT RAISE(ABORT,'events are append-only'); END;
            CREATE TABLE IF NOT EXISTS sessions (
                id INTEGER PRIMARY KEY, name TEXT UNIQUE NOT NULL, workspace TEXT NOT NULL,
                cwd TEXT NOT NULL, monty BLOB, messages TEXT NOT NULL,
                generation TEXT NOT NULL, pending TEXT, last_event_id INTEGER NOT NULL);
            CREATE TABLE IF NOT EXISTS files (
                session_id INTEGER NOT NULL, path TEXT NOT NULL,
                kind TEXT NOT NULL CHECK(kind IN ('file','directory')),
                storage_key TEXT, size INTEGER NOT NULL, generation TEXT NOT NULL,
                PRIMARY KEY (session_id,path));
        ''')
        try:
            with self._transaction():
                if session_id is None:
                    self.id = self.db.execute('SELECT COALESCE(MAX(session_id)+1,0) FROM events').fetchone()[0]
                    if not 0 <= self.id < 65536:
                        raise ValueError('names.py capacity exhausted')
                    self.name = num_to_name(self.id)
                    self.storage.create(self.name)
                    fs = Unix(cwd=self.home)
                    fs.mkdir(self.home)
                    fs.mkdir('/tmp')
                    payload = self._publish(fs.entries())
                    self._append('session_created', {**payload, 'name': self.name, 'cwd': self.home})
                else:
                    self.id = session_id
                    self._replay()
                    self.name = self.state.name
            self._checkout()
        except BaseException:
            self.db.close()
            raise
        self.model = model or os.environ.get('LLM_MODEL', 'gpt-6-luna')
        self.base_url = (base_url or os.environ.get('LLM_BASE_URL', 'https://api.openai.com/v1')).rstrip('/')
        self.api_key = api_key if api_key is not None else os.environ.get('LLM_API_KEY', '')

    @contextmanager
    def _transaction(self):
        self.db.execute('BEGIN IMMEDIATE')
        try:
            yield
            self.db.execute('COMMIT')
        except BaseException:
            if self.db.in_transaction:
                self.db.execute('ROLLBACK')
            raise

    @property
    def home(self):
        return f'/home/{self.name}'

    @property
    def state(self):
        row = self.db.execute('SELECT id,name,workspace,cwd,monty,messages,generation,pending,last_event_id '
                              'FROM sessions WHERE id=?', (self.id,)).fetchone()
        if row is None:
            raise ValueError(f'unknown session: {self.id}')
        return Session(*row[:5], json.loads(row[5]), row[6],
                       None if row[7] is None else json.loads(row[7]), row[8])

    def _append(self, kind, payload, monty=None):
        if not self.db.in_transaction:
            raise RuntimeError('events require a transaction')
        cursor = self.db.execute('INSERT INTO events(session_id,kind,execution_id,payload,monty) VALUES (?,?,?,?,?)',
                                 (self.id, kind, payload.get('execution_id'), json.dumps(payload), monty))
        self._project(cursor.lastrowid, kind, payload, monty)

    def _project(self, event_id, kind, payload, monty):
        if kind == 'session_created':
            state = Session(self.id, payload['name'], payload['name'], payload['cwd'],
                            None, [], payload['generation'], None, event_id)
        else:
            state = self.state
        if kind == 'message_appended':
            state.messages.append(payload['message'])
        elif kind == 'execution_requested':
            state.pending = payload
        elif kind in ('execution_committed', 'execution_failed'):
            state.pending = None
            if payload['tool_id'] is not None:
                state.messages.append({'role': 'tool', 'tool_call_id': payload['tool_id'], 'content': payload['output']})
            if kind == 'execution_committed':
                state.cwd, state.monty, state.generation = payload['cwd'], monty, payload['generation']
        elif kind != 'session_created':
            raise ValueError(f'unknown event kind: {kind}')
        self.db.execute('INSERT OR REPLACE INTO sessions VALUES (?,?,?,?,?,?,?,?,?)',
                        (state.id, state.name, state.workspace, state.cwd, state.monty,
                         json.dumps(state.messages), state.generation,
                         None if state.pending is None else json.dumps(state.pending), event_id))
        if kind in ('session_created', 'execution_committed'):
            self.db.execute('DELETE FROM files WHERE session_id=?', (self.id,))
            self.db.executemany('INSERT INTO files VALUES (?,?,?,?,?,?)',
                [(self.id, f['path'], f['kind'], f['storage_key'], f['size'], payload['generation'])
                 for f in payload['files']])

    def _replay(self):
        row = self.db.execute('SELECT last_event_id FROM sessions WHERE id=?', (self.id,)).fetchone()
        after = row[0] if row else 0
        for event_id, kind, payload, monty in self.db.execute(
                'SELECT id,kind,payload,monty FROM events WHERE session_id=? AND id>? ORDER BY id', (self.id, after)):
            self._project(event_id, kind, json.loads(payload), monty)

    def rebuild(self):
        """Rebuild projections and live files from events, without executing code."""
        with self._transaction():
            self.db.execute('DELETE FROM files WHERE session_id=?', (self.id,))
            self.db.execute('DELETE FROM sessions WHERE id=?', (self.id,))
            self._replay()
        self._checkout()

    def _checkout(self):
        # The directory is a cache. Interrupted copies are rebuilt on reopen/retry.
        with self._transaction():
            self.workspace = self.storage.checkout(self.name, self.files())

    def events(self):
        return [{'id': row[0], 'kind': row[1], 'payload': json.loads(row[2]), 'created_at': row[3]}
                for row in self.db.execute('SELECT id,kind,payload,created_at FROM events WHERE session_id=? ORDER BY id',
                                           (self.id,))]

    def append_message(self, message, *, expected_event=None):
        with self._transaction():
            state = self.state
            if state.pending is not None:
                raise RuntimeError('an execution is pending; resume it before appending messages')
            if expected_event is not None and state.last_event_id != expected_event:
                raise RuntimeError('session changed during the model request; resume the current turn')
            self._append('message_appended', {'message': message})

    def sessions(self):
        return [dict(zip(('id', 'name', 'workspace'), row))
                for row in self.db.execute('SELECT id,name,workspace FROM sessions ORDER BY id')]

    def _publish(self, entries):
        return {'generation': uuid.uuid4().hex, 'files': [
            {'path': path, 'kind': 'directory' if content is None else 'file',
             'storage_key': None if content is None else self.storage.put(self.name, content),
             'size': 0 if content is None else len(content)} for path, content in entries]}

    def files(self):
        return [dict(zip(('path', 'kind', 'storage_key', 'size', 'generation'), row)) for row in
                self.db.execute('SELECT path,kind,storage_key,size,generation FROM files '
                                'WHERE session_id=? ORDER BY path', (self.id,))]

    def read_file(self, path):
        path = '/' + posixpath.normpath('/' + str(path).lstrip('/')).lstrip('/')
        row = self.db.execute('SELECT kind,storage_key FROM files WHERE session_id=? AND path=?',
                              (self.id, path)).fetchone()
        if row is None:
            raise FileNotFoundError(path)
        if row[0] != 'file':
            raise IsADirectoryError(path)
        return self.storage.read(row[1])

    def close(self):
        self.db.close()

    def _emit(self, kind, text):
        if self.on_output is not None and (text or kind == 'result'):
            self.on_output(kind, text)

    def _result(self, execution_id):
        row = self.db.execute("SELECT kind,payload FROM events WHERE session_id=? AND execution_id=? "
                              "AND kind IN ('execution_committed','execution_failed')", (self.id, execution_id)).fetchone()
        if row is None:
            return None
        payload = json.loads(row[1])
        if row[0] == 'execution_failed':
            raise ExecutionError(payload['output'])
        return payload['output']

    def execute(self, code=None, *, tool_id=None, execution_id=None):
        """Resume with execute(); reuse execution_id to retrieve a committed result."""
        self._stdout = ''
        output = self._execute(code, tool_id, execution_id)
        self._checkout()
        self._emit('result', output[len(self._stdout):])
        return output

    def _execute(self, code, tool_id, execution_id):
        with self._transaction():
            pending = self.state.pending
            execution_id = execution_id or ('tool:' + tool_id if tool_id is not None else
                                            pending['execution_id'] if pending else uuid.uuid4().hex)
            requested = self.db.execute("SELECT payload FROM events WHERE session_id=? AND execution_id=? "
                                        "AND kind='execution_requested'", (self.id, execution_id)).fetchone()
            if requested:
                request = json.loads(requested[0])
                if code is not None and (request['code'] != code or request['tool_id'] != tool_id):
                    raise ValueError('execution ID was already used for another command')
            elif pending:
                raise RuntimeError('an execution is pending; call execute() or run() to resume')
            else:
                if not isinstance(code, str):
                    raise ValueError('code must be a string')
                self._append('execution_requested', {'execution_id': execution_id, 'code': code, 'tool_id': tool_id})
        # The request is now durable, independent of the transaction below.
        with self._transaction():
            result = self._result(execution_id)
            if result is not None:
                return result
            request = self.state.pending
            if request['tool_id'] not in self._streamed_calls:
                self._emit('python', request['code'])
            try:
                fs, snapshot, output = self._evaluate(request['code'])
            except MontyCrashedError:
                raise  # worker failures keep the request pending
            except MontyError as exc:
                output = f'{type(exc).__name__}: {exc}'[:24000]
                self._append('execution_failed', {**request, 'output': output})
            else:
                payload = self._publish(fs.entries())
                self._append('execution_committed', {**request, **payload, 'cwd': fs.cwd, 'output': output}, snapshot)
        return self._result(execution_id)

    def _evaluate(self, code):
        state = self.state
        entries = [(f['path'], None if f['kind'] == 'directory' else self.storage.read(f['storage_key']))
                   for f in self.files()]
        fs = Unix(entries, state.cwd, environ={
            'HOME': self.home, 'USER': self.name, 'LOGNAME': self.name,
            'PWD': state.cwd, 'TMPDIR': '/tmp', 'LANG': 'C.UTF-8',
        })
        fs.path_iterdir(Path(state.cwd))
        functions = {'pwd': fs.pwd, 'cd': fs.cd, 'ls': fs.ls, 'mkdir': fs.mkdir,
                     'read': fs.read, 'write': fs.write, 'edit': fs.edit,
                     'glob': fs.glob, 'grep': fs.grep, 'mv': fs.mv, 'rm': fs.rm,
                     'numpy': fs.numpy, 'pandas': fs.pandas, 'PIL': fs.PIL,
                     'openpyxl': fs.openpyxl, 'pypdf': fs.pypdf, 'docx': fs.docx}
        with Monty(max_processes=1, request_timeout=20) as pool:
            with pool.checkout(limits={'max_memory': 64 * 1024**2,
                                       'max_feed_duration_secs': 10, 'max_suspensions': 1000}) as session:
                if state.monty is not None:
                    session.load_session(state.monty)
                printed, byte_count = [], 0
                def capture(stream, text):
                    nonlocal byte_count
                    byte_count += len(text.encode())
                    if byte_count > 16000:
                        raise MemoryError('printed output exceeds 16000 bytes')
                    printed.append(text)
                    self._emit('stdout', text)
                value = session.feed_run(code, external_lookup=functions, os=fs,
                                         cwd=fs.cwd, print_callback=capture)
                self._stdout = ''.join(printed)
                output = (self._stdout + (repr(value) if value is not None else ''))[:24000]
                return fs, session.dump(), output

    def _tool(self, call, revision):
        try:
            if call['function']['name'] != 'python':
                raise ValueError('unknown tool')
            code = json.loads(call['function']['arguments'])['code']
            if not isinstance(code, str):
                raise ValueError('code must be a string')
        except (ValueError, KeyError, TypeError) as exc:
            self.append_message({'role': 'tool', 'tool_call_id': call['id'], 'content': str(exc)},
                                expected_event=revision)
            self._emit('error', str(exc))
            return
        try:
            self.execute(code, tool_id=call['id'])
        except ExecutionError as exc:
            self._emit('error', str(exc))  # The failure is already recorded.

    def _response(self, client, messages, schema):
        """Display deltas immediately; return only a complete assistant message."""
        text, calls, shown, finish = [], {}, {}, None
        with client.stream('POST', self.base_url + '/chat/completions',
                           headers={'Authorization': f'Bearer {self.api_key}'},
                           json={'model': self.model, 'stream': True, 'messages': [
                               {'role': 'system', 'content':
                                f'Your session ID is {self.id}; your name is {self.name}.\n' + HELP},
                               *messages], 'tools': [schema], 'parallel_tool_calls': False}) as response:
            response.raise_for_status()
            for data in sse_data(response):
                if data == '[DONE]':
                    break
                packet = json.loads(data)
                if 'error' in packet:
                    raise RuntimeError(str(packet['error']))
                for choice in packet.get('choices', []):
                    if choice.get('index', 0) != 0:
                        continue
                    delta = choice.get('delta', {})
                    content = delta.get('content') or delta.get('refusal') or ''
                    text.append(content)
                    self._emit('assistant', content)
                    for part in delta.get('tool_calls', []):
                        index = part['index']
                        call = calls.setdefault(index, {'id': '', 'type': 'function',
                                                       'function': {'name': '', 'arguments': ''}})
                        call['id'] += part.get('id', '')
                        function = part.get('function', {})
                        for field in ('name', 'arguments'):
                            call['function'][field] += function.get(field, '')
                        code = code_prefix(call['function']['arguments'])
                        self._emit('python', code[len(shown.get(index, '')):])
                        shown[index] = code
                    finish = choice.get('finish_reason') or finish
        if finish not in ('stop', 'tool_calls') or (finish == 'tool_calls' and not calls):
            raise RuntimeError('Incomplete model response; resume to retry (partial code was not executed)')
        if any(not call['id'] or not call['function']['name'] for call in calls.values()):
            raise RuntimeError('Incomplete tool call; resume to retry')
        message = {'role': 'assistant', 'content': ''.join(text)}
        if calls:
            message['tool_calls'] = [calls[index] for index in sorted(calls)]
            self._streamed_calls.update(call['id'] for call in calls.values())
        return message

    def run(self, prompt=None, *, max_steps=20):
        """Resume pending work automatically before accepting a new user turn."""
        if self.state.pending:
            try:
                self.execute()
            except ExecutionError as exc:
                self._emit('error', str(exc))
        messages = self.state.messages
        if prompt is not None:
            if messages and (messages[-1]['role'] in ('user', 'tool') or messages[-1].get('tool_calls')):
                repeated = messages[-1]['role'] == 'user' and messages[-1]['content'] == prompt
                answer = self.run(max_steps=max_steps)
                if repeated:
                    return answer
            self.append_message({'role': 'user', 'content': prompt})
        if not self.state.messages:
            return ''
        schema = {'type': 'function', 'function': {'name': 'python',
                  'description': 'Run Python in the durable Monty virtual Unix environment.',
                  'parameters': {'type': 'object', 'properties': {'code': {'type': 'string'}},
                                 'required': ['code'], 'additionalProperties': False}}}
        with httpx.Client(timeout=120) as client:
            for _ in range(max_steps):
                state = self.state
                messages = state.messages
                assistant = next((m for m in reversed(messages) if m['role'] == 'assistant'), {})
                answered = {m['tool_call_id'] for m in messages if m['role'] == 'tool'}
                call = next((c for c in assistant.get('tool_calls', []) if c['id'] not in answered), None)
                if call:
                    self._tool(call, state.last_event_id)
                    continue
                if messages[-1]['role'] == 'assistant':
                    return messages[-1].get('content') or ''
                message = self._response(client, messages, schema)
                self.append_message(message, expected_event=state.last_event_id)
                if message['role'] == 'assistant' and not message.get('tool_calls'):
                    return message.get('content') or ''
        raise RuntimeError('Step budget reached; call run() to continue the durable turn')
